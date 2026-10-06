"""Leitura das planilhas desta pasta e checagem das conexões entre as bases.

    bases.xlsx         pacientes, medicos, consultas, horarios, exames  (dados estruturados)
    conhecimento.xlsx  documentos                                       (textos para o RAG)

Cada aba vira uma lista de dicionários. Abas que começam com "_" são só explicação e ficam de fora.
"""
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

PASTA_DADOS = Path(__file__).resolve().parent

# Como as bases se conectam: (base, coluna) -> (base de destino, chave)
CONEXOES = {
    ("consultas", "paciente_id"): ("pacientes", "paciente_id"),
    ("consultas", "medico_id"): ("medicos", "medico_id"),
    ("horarios", "medico_id"): ("medicos", "medico_id"),
    ("exames", "paciente_id"): ("pacientes", "paciente_id"),
    ("exames", "medico_id"): ("medicos", "medico_id"),
}


def _valor(v):
    """Normaliza o que vem do Excel para tipos que viram JSON sem surpresa."""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, str):
        return v.strip()
    return v


def ler_planilha(caminho: Path) -> dict[str, list[dict]]:
    """Uma aba = uma base. A primeira linha é o cabeçalho."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    bases = {}
    for ws in wb.worksheets:
        if ws.title.startswith("_"):
            continue
        linhas = ws.iter_rows(values_only=True)
        cabecalho = [str(c).strip() for c in next(linhas)]
        bases[ws.title] = [
            dict(zip(cabecalho, map(_valor, linha)))
            for linha in linhas
            if any(v is not None for v in linha)
        ]
    return bases


def carregar_bases() -> dict[str, list[dict]]:
    return ler_planilha(PASTA_DADOS / "bases.xlsx")


def carregar_documentos() -> list[dict]:
    return ler_planilha(PASTA_DADOS / "conhecimento.xlsx")["documentos"]


# Colunas com data (AAAA-MM-DD) em cada base: são elas que andam junto quando a agenda é atualizada
COLUNAS_DE_DATA = {"consultas": "data", "horarios": "data", "exames": "data_pedido"}


def manter_no_futuro(bases: dict[str, list[dict]], hoje: date | None = None) -> int:
    """Avança todas as datas em semanas inteiras, até a primeira consulta agendada não estar no passado.

    As datas do Excel são fixas. Sem isso, num dia de aula mais adiante a agenda estaria vencida e o agente
    não acharia horário livre. Andar em semanas mantém os dias da semana (a consulta de quinta continua na quinta).
    Altera as bases e devolve quantas semanas avançou (0 se já estava em dia).
    """
    hoje = hoje or date.today()
    agendadas = [date.fromisoformat(c["data"]) for c in bases["consultas"] if c["status"] == "agendada"]
    if not agendadas or min(agendadas) >= hoje:
        return 0
    semanas = -(-(hoje - min(agendadas)).days // 7)  # arredonda para cima
    for base, coluna in COLUNAS_DE_DATA.items():
        for linha in bases[base]:
            linha[coluna] = (date.fromisoformat(linha[coluna]) + timedelta(weeks=semanas)).isoformat()
    return semanas


def validar_conexoes(bases: dict[str, list[dict]]) -> list[str]:
    """Devolve a lista de problemas (vazia quando toda chave estrangeira existe na base de destino)."""
    problemas = []
    for (origem, coluna), (destino, chave) in CONEXOES.items():
        existentes = {linha[chave] for linha in bases[destino]}
        for linha in bases[origem]:
            if linha[coluna] not in existentes:
                problemas.append(f"{origem}.{coluna}={linha[coluna]} não existe em {destino}")
    return problemas
